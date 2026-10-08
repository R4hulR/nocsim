"""Build an editable Word version (main.docx) of the paper from main.tex.

Pandoc does not understand the IEEEtran-specific commands, and on its own it
would lose cross-reference and citation numbers. This script:
  1. reads the label/citation numbers LaTeX itself assigned (from main.aux),
  2. writes a simplified copy of main.tex with every \\ref, \\cite, section
     heading and caption replaced by the exact numbers used in the PDF,
  3. converts that copy with pandoc (equations become native Word equations),
  4. sets a plain journal-style look (Times New Roman, justified text).

Usage:  python make_docx.py path/to/main.aux
(produce main.aux with:  tectonic -X compile main.tex --keep-intermediates)
Requires: pypandoc_binary, python-docx
"""

import re
import sys
from pathlib import Path

import pypandoc
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt

HERE = Path(__file__).resolve().parent


def read_aux(aux: Path) -> tuple[dict, dict]:
    labels, cites = {}, {}
    for line in aux.read_text(encoding="utf-8").splitlines():
        m = re.match(r"\\newlabel\{([^}]*)\}\{\{(.*?)\}\{", line)
        if m:
            labels[m.group(1)] = re.sub(r"\\mbox\s*\{(.*?)\}", r"\1", m.group(2)).strip()
        m = re.match(r"\\bibcite\{([^}]*)\}\{(\d+)\}", line)
        if m:
            cites[m.group(1)] = int(m.group(2))
    return labels, cites


def roman(n: int) -> str:
    vals = [(10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")]
    out = ""
    for v, s in vals:
        while n >= v:
            out, n = out + s, n - v
    return out


def simplify(tex: str, labels: dict, cites: dict) -> str:
    # --- front matter: plain author block, title footnote as a normal line
    tex = tex.replace("\\IEEEoverridecommandlockouts % allows \\thanks in conference mode\n", "")
    preprint = ""
    start = tex.find("%\n\\thanks{")
    if start >= 0:
        # Find the matching closing brace of \thanks{...} (it contains nested braces).
        i = tex.index("{", start) + 1
        depth = 1
        while depth:
            depth += {"{": 1, "}": -1}.get(tex[i], 0)
            i += 1
        preprint = tex[tex.index("{", start) + 1 : i - 1]
        tex = tex[:start] + tex[i:]
    tex = re.sub(r"\\author\{.*?\n\n", "", tex, count=1, flags=re.S)
    front = (
        "\\begin{center}\nRahul Ray\\\\\nB.Tech.\\ CSE (2023), Assam University, Silchar, India\\\\\n"
        "rahulrayconnect@gmail.com\\\\\n" + preprint.replace("Preprint.", "Preprint:") + "\n\\end{center}\n\n"
    )
    # Pandoc always hoists an abstract environment above everything else, so
    # make it an ordinary heading that follows the author block.
    tex = tex.replace("\\begin{abstract}\n", front + "\\section*{Abstract}\n", 1)
    tex = tex.replace("\\end{abstract}", "", 1)
    tex = tex.replace("\\begin{IEEEkeywords}\n", "\\noindent\\textit{Index Terms}: ").replace("\\end{IEEEkeywords}", "")

    # --- numbered headings, as in the IEEE PDF (I., II., ... and A., B., ...)
    sec = sub = 0
    def heading(m):
        nonlocal sec, sub
        kind, star, title = m.group(1), m.group(2), m.group(3)
        if star:
            return m.group(0)
        if kind == "section":
            sec, sub = sec + 1, 0
            return f"\\section*{{{roman(sec)}. {title}}}"
        sub += 1
        return f"\\subsection*{{{chr(64 + sub)}. {title}}}"
    tex = re.sub(r"\\(section|subsection)(\*?)\{([^}]*)\}", heading, tex)

    # --- captions: "Fig. n." and "TABLE n." prefixes in document order
    def number_floats(env: str, prefix_fn):
        out, pos = [], 0
        for m in re.finditer(r"\\begin\{" + env + r"\*?\}.*?\\end\{" + env + r"\*?\}", tex, re.S):
            block = m.group(0)
            lab = re.search(r"\\label\{([^}]*)\}", block)
            num = labels.get(lab.group(1), "?") if lab else "?"
            block = block.replace("\\caption{", "\\caption{" + prefix_fn(num), 1)
            out.append(tex[pos:m.start()] + block)
            pos = m.end()
        return "".join(out) + tex[pos:]
    tex = number_floats("figure", lambda n: f"Fig. {n}. ")
    tex = number_floats("table", lambda n: f"TABLE {n}. ")
    tex = tex.replace("\\begin{figure*}", "\\begin{figure}").replace("\\end{figure*}", "\\end{figure}")

    # --- cross-references and citations -> literal numbers
    tex = re.sub(r"\\ref\{([^}]*)\}", lambda m: labels.get(m.group(1), "??"), tex)
    tex = re.sub(r"~?\\cite\{([^}]*)\}",
                 lambda m: " [" + ", ".join(str(cites[k.strip()]) for k in m.group(1).split(",")) + "]", tex)

    # --- grid sizes like $8\times8$ as plain text ("8×8"), not Word equations,
    # so submission systems that scrape the abstract get clean text
    tex = re.sub(r"\$(\d+|n)\s*\\times\s*(\d+|n)\$", lambda m: f"{m.group(1)}×{m.group(2)}", tex)

    # --- tables: pandoc can't do \multirow / \cmidrule
    tex = re.sub(r"\\multirow\{\d+\}\{\*\}\{([^}]*)\}", r"\1", tex)
    tex = re.sub(r"\\cmidrule(\([a-z]*\))?\{[^}]*\}", "", tex)

    # --- bibliography -> plain numbered paragraphs
    def bib(m):
        body = m.group(1)
        items = re.split(r"\\bibitem\{([^}]*)\}", body)[1:]
        paras = [f"[{cites[items[i]]}] {' '.join(items[i + 1].split())}" for i in range(0, len(items), 2)]
        return "\\section*{References}\n\n" + "\n\n".join(paras)
    tex = re.sub(r"\\begin\{thebibliography\}\{\d+\}(.*?)\\end\{thebibliography\}", bib, tex, flags=re.S)
    return tex


def _set_font(style_obj, name: str) -> None:
    """Set a style's font, removing theme-font links that would override it."""
    style_obj.font.name = name
    rpr = style_obj.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            rfonts.attrib.pop(qn(attr), None)
        for attr in ("w:ascii", "w:hAnsi", "w:cs"):
            rfonts.set(qn(attr), name)


def _grid(table) -> None:
    """Thin single-line borders on every cell."""
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "808080")
        borders.append(el)
    tbl_pr.append(borders)


def style(docx_path: Path) -> None:
    doc = Document(str(docx_path))
    names = {s.name for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Abstract", "Caption",
                 "Image Caption", "Table Caption", "Title", "Heading 1", "Heading 2", "Heading 3"):
        if name not in names:
            continue
        st = doc.styles[name]
        _set_font(st, "Times New Roman")
        if name in ("Normal", "Body Text", "First Paragraph", "Compact", "Abstract"):
            st.font.size = Pt(11)
        elif name in ("Caption", "Image Caption", "Table Caption"):
            st.font.size = Pt(9.5)
        elif name == "Title":
            st.font.size = Pt(18)
        elif name.startswith("Heading"):
            st.font.size = Pt(12)
            st.font.color.rgb = None
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.style.name in ("Body Text", "First Paragraph"):
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if p.text.startswith("Rahul Ray"):  # author block: centred, not justified
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for p in paras:
        if p.style.name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Every figure at the full text width, keeping its aspect ratio.
    # Pandoc leaves the page size unset; use A4 with 1-inch margins.
    for section in doc.sections:
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
    section = doc.sections[0]
    text_width = section.page_width - section.left_margin - section.right_margin
    for shape in doc.inline_shapes:
        ratio = shape.height / shape.width
        shape.width = int(text_width)
        shape.height = int(text_width * ratio)

    for t in doc.tables:
        _grid(t)
    doc.save(str(docx_path))


def main():
    aux = Path(sys.argv[1]) if len(sys.argv) > 1 else HERE / "main.aux"
    labels, cites = read_aux(aux)
    tex = simplify((HERE / "main.tex").read_text(encoding="utf-8"), labels, cites)
    tmp = HERE / "_pandoc_main.tex"
    tmp.write_text(tex, encoding="utf-8")
    out = HERE / "main.docx"
    try:
        pypandoc.convert_file(str(tmp), "docx", format="latex", outputfile=str(out),
                              extra_args=[f"--resource-path={HERE / 'figures'}"])
    finally:
        tmp.unlink()
    style(out)
    print("wrote", out)


if __name__ == "__main__":
    main()
